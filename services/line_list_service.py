# -*- coding: utf-8 -*-
"""
services/line_list_service.py – PipeAgent
سرویس جامع مدیریت، اعتبارسنجی مهندسی و ایمپورت/اکسپورت لیست خطوط لوله‌کشی (Piping Line List)
منطبق با استانداردهای ASME B31.3، دارای موتور پارس هوشمند تگ نامبر،
محاسبه اتوماتیک فشار هیدروتست، ایمپورت دسته‌ای سریع و داشبورد شاخص‌های خطوط پروژه.
"""

from __future__ import annotations

import csv
import logging
import re
from datetime import datetime
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple, Union

from sqlalchemy import func, and_, or_, desc, asc, case
from sqlalchemy.orm import Session
from sqlalchemy.exc import SQLAlchemyError

from config import FLUID_SERVICES, NDT_EXTENT_BY_SERVICE
from db.manager import DatabaseManager
from db.models import LineListItem, Project
from services.license import increment_usage

logger = logging.getLogger(__name__)


# ──────────────────────────────────────────────
#  Enums & ASME B31.3 Fluid Services
# ──────────────────────────────────────────────

class FluidCategory(str, Enum):
    """دسته‌بندی سرویس‌های سیال طبق استاندارد ASME B31.3"""
    NORMAL_FLUID_SERVICE = "Normal Fluid Service"      # Implementation note.
    CATEGORY_D = "Category D"                          # Implementation note.
    CATEGORY_M = "Category M"                          # Implementation note.
    HIGH_PRESSURE = "High Pressure (Chapter IX)"       # Implementation note.
    HIGH_PURITY = "High Purity (Chapter X)"            # Implementation note.
    SEVERE_CYCLIC = "Severe Cyclic Conditions"         # Implementation note.


class TestMedium(str, Enum):
    POTABLE_WATER = "Water"
    DEMIN_WATER = "Demin Water"
    AIR = "Air"
    NITROGEN = "Nitrogen"
    SERVICE_FLUID = "Service Fluid"


# Implementation note.
COLUMN_MAP: Dict[str, str] = {
    # Implementation note.
    "line_number": "line_number", "line no": "line_number", "line no.": "line_number",
    "lineno": "line_number", "line": "line_number", "piping line number": "line_number",
    "line tag": "line_number", "tag number": "line_number",
    # Implementation note.
    "pid": "pid_number", "p&id": "pid_number", "pid_number": "pid_number",
    "p&id number": "pid_number", "pid no": "pid_number", "pid dwg": "pid_number",
    "iso": "iso_number", "iso_number": "iso_number", "isometric": "iso_number",
    "iso no": "iso_number", "iso dwg": "iso_number",
    # Implementation note.
    "fluid_code": "fluid_code", "fluid code": "fluid_code", "fluid": "fluid_code",
    "fluid_name": "fluid_name", "fluid name": "fluid_name", "service": "fluid_name",
    "fluid_service": "fluid_service", "fluid service": "fluid_service", "b31.3 service": "fluid_service",
    # Implementation note.
    "design_pressure": "design_pressure_barg", "design pressure": "design_pressure_barg",
    "des. press": "design_pressure_barg", "dp": "design_pressure_barg", "design press (barg)": "design_pressure_barg",
    "design_temp": "design_temp_c", "design temperature": "design_temp_c",
    "des. temp": "design_temp_c", "dt": "design_temp_c", "design temp (c)": "design_temp_c",
    "operating_pressure": "operating_pressure_barg", "op. press": "operating_pressure_barg", "op": "operating_pressure_barg",
    "operating_temp": "operating_temp_c", "op. temp": "operating_temp_c", "ot": "operating_temp_c",
    "test_pressure": "test_pressure_barg", "test pressure": "test_pressure_barg", "tp": "test_pressure_barg",
    "hydro test pressure": "test_pressure_barg", "test press (barg)": "test_pressure_barg",
    "test_medium": "test_medium", "test medium": "test_medium", "testing fluid": "test_medium",
    # Implementation note.
    "pipe_class": "pipe_class", "pipe class": "pipe_class", "class": "pipe_class", "piping class": "pipe_class",
    "pipe_spec": "pipe_spec", "spec": "pipe_spec", "piping spec": "pipe_spec",
    "size": "size_nps", "size_nps": "size_nps", "nps": "size_nps", "dn": "size_nps", "nominal size": "size_nps",
    "schedule": "schedule", "sch": "schedule", "wall thickness": "schedule",
    "material": "material", "mat": "material", "matl": "material", "material grade": "material",
    "insulation": "insulation", "insul": "insulation", "insulation type": "insulation", "insulation thk": "insulation",
    "tracing": "tracing", "heat tracing": "tracing",
    "painting": "painting_code", "paint": "painting_code", "painting system": "painting_code", "paint spec": "painting_code",
    # Implementation note.
    "from": "from_point", "from_point": "from_point", "from point": "from_point", "origin": "from_point",
    "to": "to_point", "to_point": "to_point", "to point": "to_point", "destination": "to_point",
    # Implementation note.
    "rt%": "ndt_percent_rt", "rt": "ndt_percent_rt", "rt extent": "ndt_percent_rt",
    "ut%": "ndt_percent_ut", "ut": "ndt_percent_ut", "ut extent": "ndt_percent_ut",
    "pwht": "pwht_required", "pwht_required": "pwht_required", "heat treatment": "pwht_required",
    # Implementation note.
    "area": "area_name", "unit": "area_name", "plant area": "area_name",
    "subsystem": "subsystem", "sub-system": "subsystem",
    "status": "status", "remarks": "remarks", "remark": "remarks", "notes": "remarks",
}


# ──────────────────────────────────────────────
#  Formatting & Engineering Conversion Helpers
# ──────────────────────────────────────────────

def _normalize_header(h: Any) -> str:
    """استانداردسازی رشته‌های هدر جهت تطابق بدون حساسیت به حروف، فاصله و خط تیره"""
    return str(h or "").strip().lower().replace("_", " ").replace("-", " ").replace(".", "")


def _to_float(val: Any) -> Optional[float]:
    """تبدیل ایمن مقادیر اعشاری مهندسی با حذف کاما و فاصله‌ها"""
    if val in (None, ""):
        return None
    try:
        clean_str = str(val).replace(",", "").replace("barg", "").replace("°C", "").replace("C", "").strip()
        return float(clean_str)
    except (ValueError, TypeError):
        return None


def _to_bool(val: Any) -> bool:
    """تبدیل مقادیر بولی برای فیلدهایی نظیر الزامات PWHT و عایق"""
    if val in (None, ""):
        return False
    s = str(val).strip().lower()
    return s in ("1", "true", "yes", "y", "x", "pwht", "req", "required", "دارد")


# ──────────────────────────────────────────────
#  LineListService Implementation
# ──────────────────────────────────────────────

class LineListService:
    """
    سرویس مرکزی مدیریت، اعتبارسنجی و ایمپورت/اکسپورت لیست خطوط لوله‌کشی پروژه
    """

    def __init__(self, db: DatabaseManager):
        self.db = db

    # Implementation note.

    def add_line(
        self,
        project_id: int,
        line_number: str,
        auto_calculate_test_pressure: bool = True,
        **kwargs,
    ) -> Optional[LineListItem]:
        """
        ثبت رسمی خط لوله جدید با اعمال اعتبارسنجی‌های استاندارد ASME B31.3
        و پر کردن خودکار فشار هیدروتست و درصد NDT در صورت عدم ورود.
        """
        clean_line_no = str(line_number).strip().upper()
        if not clean_line_no:
            return None

        if not increment_usage(self.db):
            logger.error("License limit reached. Cannot create new Line List item.")
            return None

        with self.db.session_scope() as session:
            existing = (
                session.query(LineListItem)
                .filter(
                    LineListItem.project_id == project_id,
                    LineListItem.line_number == clean_line_no,
                )
                .first()
            )
            if existing:
                logger.warning(f"Line '{clean_line_no}' already exists in Project #{project_id}.")
                return None

            # Implementation note.
            fluid_service = kwargs.get("fluid_service") or FluidCategory.NORMAL_FLUID_SERVICE.value
            extent = NDT_EXTENT_BY_SERVICE.get(fluid_service, {})

            rt_percent = _to_float(kwargs.get("ndt_percent_rt"))
            if rt_percent is None:
                rt_percent = extent.get("RT", 5.0)  # Implementation note.

            ut_percent = _to_float(kwargs.get("ndt_percent_ut"))
            if ut_percent is None:
                ut_percent = extent.get("UT", 0.0)

            # Implementation note.
            design_p = _to_float(kwargs.get("design_pressure_barg"))
            operating_p = _to_float(kwargs.get("operating_pressure_barg"))
            test_p = _to_float(kwargs.get("test_pressure_barg"))

            # Implementation note.
            if test_p is None and design_p is not None and auto_calculate_test_pressure:
                test_p = round(design_p * 1.5, 2)
                logger.info(f"Auto-calculated hydrotest pressure for '{clean_line_no}': {test_p} barg (1.5x DP).")

            # Implementation note.
            parsed_data = self.parse_line_tag(clean_line_no)

            item = LineListItem(
                project_id=project_id,
                line_number=clean_line_no,
                pid_number=kwargs.get("pid_number"),
                iso_number=kwargs.get("iso_number"),
                fluid_code=kwargs.get("fluid_code") or parsed_data.get("fluid_code"),
                fluid_name=kwargs.get("fluid_name"),
                fluid_service=fluid_service,
                design_pressure_barg=design_p,
                design_temp_c=_to_float(kwargs.get("design_temp_c")),
                operating_pressure_barg=operating_p,
                operating_temp_c=_to_float(kwargs.get("operating_temp_c")),
                test_pressure_barg=test_p,
                test_medium=kwargs.get("test_medium") or TestMedium.POTABLE_WATER.value,
                pipe_class=kwargs.get("pipe_class") or parsed_data.get("pipe_class"),
                pipe_spec=kwargs.get("pipe_spec"),
                size_nps=kwargs.get("size_nps") or parsed_data.get("size_nps"),
                schedule=kwargs.get("schedule"),
                material=kwargs.get("material"),
                insulation=kwargs.get("insulation") or parsed_data.get("insulation"),
                tracing=kwargs.get("tracing") or parsed_data.get("tracing"),
                painting_code=kwargs.get("painting_code"),
                ndt_percent_rt=rt_percent,
                ndt_percent_ut=ut_percent,
                pwht_required=_to_bool(kwargs.get("pwht_required")),
                from_point=kwargs.get("from_point"),
                to_point=kwargs.get("to_point"),
                area_id=kwargs.get("area_id"),
                subsystem=kwargs.get("subsystem"),
                status=kwargs.get("status") or "Active",
                remarks=kwargs.get("remarks"),
                created_at=datetime.utcnow(),
            )
            session.add(item)
            session.flush()

            logger.info(f"Line '{item.line_number}' registered successfully in Project #{project_id}.")
            return item

    # Implementation note.

    @staticmethod
    def parse_line_tag(line_tag: str) -> Dict[str, Optional[str]]:
        """
        استخراج هوشمند اجزای ساختار شماره‌گذاری خطوط پایپینگ بر اساس استانداردهای متداول پالایشگاهی:
        مثال: '10"-HC-1002-1CS1P01-H50-ET'
        -> سایز: 10" | سیال: HC | شماره: 1002 | کلاس: 1CS1P01 | عایق: H50 | هیت‌تریس: ET
        """
        tag = line_tag.strip().upper()
        result: Dict[str, Optional[str]] = {
            "size_nps": None,
            "fluid_code": None,
            "line_sequence": None,
            "pipe_class": None,
            "insulation": None,
            "tracing": None,
        }

        # Implementation note.
        pattern = re.compile(
            r"^(?P<size>\d+\"?|\d+\.\d+\"?|DN\d+)"
            r"[-_\s]+(?P<fluid>[A-Z0-9]+)"
            r"[-_\s]+(?P<seq>[A-Z0-9]+)"
            r"[-_\s]+(?P<class>[A-Z0-9]+)"
            r"(?:[-_\s]+(?P<ins>[A-Z0-9]+))?"
            r"(?:[-_\s]+(?P<trace>[A-Z0-9]+))?",
            re.IGNORECASE,
        )

        match = pattern.match(tag)
        if match:
            gd = match.groupdict()
            result["size_nps"] = gd.get("size")
            result["fluid_code"] = gd.get("fluid")
            result["line_sequence"] = gd.get("seq")
            result["pipe_class"] = gd.get("class")
            result["insulation"] = gd.get("ins")
            result["tracing"] = gd.get("trace")

        return result

    # Implementation note.

    def import_lines_batch(
        self,
        project_id: int,
        lines_data: List[Dict[str, Any]],
        update_existing: bool = True,
    ) -> Dict[str, Any]:
        """
        موتور فوق‌سریع و یکپارچه واردات دسته‌ای خطوط لوله در یک تراکنش اتمیک
        (حل مشکل Session Thrashing و کوئری‌های مکرر).
        """
        if not lines_data:
            return {"created": 0, "updated": 0, "skipped": 0, "errors": ["No rows provided."]}

        created_count = 0
        updated_count = 0
        skipped_count = 0
        errors: List[str] = []

        with self.db.session_scope() as session:
            # Implementation note.
            existing_lines_map = {
                l.line_number.strip().upper(): l
                for l in session.query(LineListItem).filter(LineListItem.project_id == project_id).all()
            }

            for row_idx, row in enumerate(lines_data, start=2):
                line_no = str(row.get("line_number") or "").strip().upper()
                if not line_no:
                    skipped_count += 1
                    errors.append(f"Row {row_idx}: Line number is missing – skipped.")
                    continue

                # Implementation note.
                design_p = _to_float(row.get("design_pressure_barg"))
                test_p = _to_float(row.get("test_pressure_barg"))
                if test_p is None and design_p is not None:
                    test_p = round(design_p * 1.5, 2)  # Implementation note.

                fluid_service = row.get("fluid_service") or FluidCategory.NORMAL_FLUID_SERVICE.value
                extent = NDT_EXTENT_BY_SERVICE.get(fluid_service, {})

                rt_p = _to_float(row.get("ndt_percent_rt"))
                if rt_p is None:
                    rt_p = extent.get("RT", 5.0)

                ut_p = _to_float(row.get("ndt_percent_ut"))
                if ut_p is None:
                    ut_p = extent.get("UT", 0.0)

                parsed_tag = self.parse_line_tag(line_no)

                fields_to_apply = {
                    "pid_number": row.get("pid_number"),
                    "iso_number": row.get("iso_number"),
                    "fluid_code": row.get("fluid_code") or parsed_tag.get("fluid_code"),
                    "fluid_name": row.get("fluid_name"),
                    "fluid_service": fluid_service,
                    "design_pressure_barg": design_p,
                    "design_temp_c": _to_float(row.get("design_temp_c")),
                    "operating_pressure_barg": _to_float(row.get("operating_pressure_barg")),
                    "operating_temp_c": _to_float(row.get("operating_temp_c")),
                    "test_pressure_barg": test_p,
                    "test_medium": row.get("test_medium") or TestMedium.POTABLE_WATER.value,
                    "pipe_class": row.get("pipe_class") or parsed_tag.get("pipe_class"),
                    "pipe_spec": row.get("pipe_spec"),
                    "size_nps": row.get("size_nps") or parsed_tag.get("size_nps"),
                    "schedule": row.get("schedule"),
                    "material": row.get("material"),
                    "insulation": row.get("insulation") or parsed_tag.get("insulation"),
                    "tracing": row.get("tracing") or parsed_tag.get("tracing"),
                    "painting_code": row.get("painting_code"),
                    "ndt_percent_rt": rt_p,
                    "ndt_percent_ut": ut_p,
                    "pwht_required": _to_bool(row.get("pwht_required")),
                    "from_point": row.get("from_point"),
                    "to_point": row.get("to_point"),
                    "area_id": row.get("area_id"),
                    "subsystem": row.get("subsystem"),
                    "status": row.get("status") or "Active",
                    "remarks": row.get("remarks"),
                }

                # Implementation note.
                if line_no in existing_lines_map:
                    if update_existing:
                        existing_obj = existing_lines_map[line_no]
                        for k, v in fields_to_apply.items():
                            if v is not None:
                                setattr(existing_obj, k, v)
                        existing_obj.updated_at = datetime.utcnow()
                        updated_count += 1
                    else:
                        skipped_count += 1
                        errors.append(f"Row {row_idx}: Line '{line_no}' already exists – skipped.")
                else:
                    new_item = LineListItem(
                        project_id=project_id,
                        line_number=line_no,
                        created_at=datetime.utcnow(),
                        **fields_to_apply,
                    )
                    session.add(new_item)
                    existing_lines_map[line_no] = new_item  # Implementation note.
                    created_count += 1

            session.flush()

        logger.info(f"Line List batch import complete for Project #{project_id}: Created {created_count}, Updated {updated_count}, Skipped {skipped_count}.")
        return {
            "created": created_count,
            "updated": updated_count,
            "skipped": skipped_count,
            "errors": errors[:100],
        }

    # Implementation note.

    def import_from_csv(
        self,
        project_id: int,
        file_path: Union[str, Path],
        update_existing: bool = True,
    ) -> Tuple[int, int, List[str]]:
        """ایمپورت لیست خطوط از فایل CSV با تشخیص خودکار Delimiter و انکودینگ UTF-8-SIG"""
        path = Path(file_path)
        if not path.is_file():
            return 0, 0, [f"File not found: {file_path}"]

        try:
            with open(path, "r", encoding="utf-8-sig", newline="") as f:
                sample = f.read(4096)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.DictReader(f, dialect=dialect)
                if not reader.fieldnames:
                    return 0, 0, ["CSV file has no valid header row."]

                # Implementation note.
                header_map: Dict[str, str] = {}
                for raw_header in reader.fieldnames:
                    norm = _normalize_header(raw_header)
                    mapped = COLUMN_MAP.get(norm)
                    if mapped:
                        header_map[raw_header] = mapped

                if "line_number" not in header_map.values():
                    return 0, 0, ["Mandatory 'Line Number' column not found in CSV header."]

                rows_to_process = []
                for row in reader:
                    mapped_row = {field: row[raw_col] for raw_col, field in header_map.items() if row.get(raw_col)}
                    rows_to_process.append(mapped_row)

                res = self.import_lines_batch(project_id, rows_to_process, update_existing)
                return res["created"], res["updated"], res["errors"]

        except Exception as e:
            logger.error(f"CSV import error: {e}")
            return 0, 0, [f"Import error: {str(e)}"]

    def import_from_excel(
        self,
        project_id: int,
        file_path: Union[str, Path],
        sheet_name: Optional[str] = None,
        update_existing: bool = True,
    ) -> Tuple[int, int, List[str]]:
        """ایمپورت لیست خطوط از فایل اکسل با استفاده از openpyxl و مصرف بهینه RAM"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            return 0, 0, ["openpyxl library is required for Excel import. Install via: pip install openpyxl"]

        path = Path(file_path)
        if not path.is_file():
            return 0, 0, [f"File not found: {file_path}"]

        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

            rows_gen = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_gen)
            except StopIteration:
                return 0, 0, ["Excel sheet is completely empty."]

            header_map: Dict[int, str] = {}
            for idx, cell in enumerate(header_row):
                if cell is None:
                    continue
                norm = _normalize_header(str(cell))
                mapped = COLUMN_MAP.get(norm)
                if mapped:
                    header_map[idx] = mapped

            if "line_number" not in header_map.values():
                return 0, 0, ["Could not find a 'Line Number' column in Excel header."]

            rows_to_process = []
            for row in rows_gen:
                if not any(cell not in (None, "") for cell in row):
                    continue
                mapped_row = {}
                for idx, field in header_map.items():
                    if idx < len(row) and row[idx] is not None:
                        val = str(row[idx]).strip()
                        if val:
                            mapped_row[field] = val
                rows_to_process.append(mapped_row)

            wb.close()
            res = self.import_lines_batch(project_id, rows_to_process, update_existing)
            return res["created"], res["updated"], res["errors"]

        except Exception as e:
            logger.error(f"Excel import error: {e}")
            return 0, 0, [f"Excel import error: {str(e)}"]

    # Implementation note.

    def search_lines(
        self,
        project_id: int,
        keyword: Optional[str] = None,
        fluid_code: Optional[str] = None,
        pipe_class: Optional[str] = None,
        subsystem: Optional[str] = None,
        pwht_only: Optional[bool] = None,
        high_pressure_only: Optional[bool] = None,
        page: int = 1,
        per_page: int = 50,
    ) -> Tuple[List[LineListItem], int]:
        """جستجوی پیشرفته چندمعیاره با فیلترهای مهندسی و صفحه‌بندی"""
        with self.db.session_scope() as session:
            query = session.query(LineListItem).filter(LineListItem.project_id == project_id)

            if keyword:
                search_str = f"%{keyword.strip()}%"
                query = query.filter(
                    or_(
                        LineListItem.line_number.ilike(search_str),
                        LineListItem.pid_number.ilike(search_str),
                        LineListItem.iso_number.ilike(search_str),
                        LineListItem.fluid_name.ilike(search_str),
                    )
                )

            if fluid_code:
                query = query.filter(LineListItem.fluid_code == fluid_code.strip().upper())

            if pipe_class:
                query = query.filter(LineListItem.pipe_class == pipe_class.strip().upper())

            if subsystem:
                query = query.filter(LineListItem.subsystem == subsystem.strip().upper())

            if pwht_only:
                query = query.filter(LineListItem.pwht_required == True)

            if high_pressure_only:
                query = query.filter(LineListItem.design_pressure_barg >= 50.0)

            total_count = query.count()
            items = (
                query.order_by(LineListItem.line_number.asc())
                .offset((page - 1) * per_page)
                .limit(per_page)
                .all()
            )
            return items, total_count

    # Implementation note.

    def get_line_list_kpis(self, project_id: int) -> Dict[str, Any]:
        """
        محاسبه ماتریس آماری لیست خطوط با استفاده از کوئری‌های تجمعی SQL
        """
        with self.db.session_scope() as session:
            stats = session.query(
                func.count(LineListItem.id).label("total_lines"),
                func.sum(case((LineListItem.pwht_required == True, 1), else_=0)).label("pwht_lines"),
                func.sum(case((LineListItem.design_pressure_barg >= 50.0, 1), else_=0)).label("high_pressure_lines"),
                func.sum(case((LineListItem.design_temp_c >= 200.0, 1), else_=0)).label("high_temp_lines"),
                func.sum(case((LineListItem.design_temp_c < -29.0, 1), else_=0)).label("cryogenic_lines"),
                func.sum(case((LineListItem.insulation.isnot(None), 1), else_=0)).label("insulated_lines"),
            ).filter(LineListItem.project_id == project_id).first()

            # Implementation note.
            fluid_counts = (
                session.query(LineListItem.fluid_service, func.count(LineListItem.id))
                .filter(LineListItem.project_id == project_id)
                .group_by(LineListItem.fluid_service)
                .all()
            )

            total = stats.total_lines or 0
            return {
                "project_id": project_id,
                "total_lines_count": total,
                "critical_metrics": {
                    "pwht_required_lines": stats.pwht_lines or 0,
                    "high_pressure_lines_50bar_plus": stats.high_pressure_lines or 0,
                    "high_temperature_lines_200c_plus": stats.high_temp_lines or 0,
                    "cryogenic_lines_below_minus_29c": stats.cryogenic_lines or 0,
                    "insulated_lines_count": stats.insulated_lines or 0,
                },
                "breakdown_by_fluid_service": {svc or "Unclassified": count for svc, count in fluid_counts},
            }

    # Implementation note.

    def export_to_excel(self, project_id: int, output_path: Union[str, Path]) -> str:
        """تولید خروجی اکسل با فرمت‌بندی مهندسی و رنگ‌بندی استاندارد"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise ImportError("openpyxl library is required for Excel export.")

        path_obj = Path(output_path)
        path_obj.parent.mkdir(parents=True, exist_ok=True)

        with self.db.session_scope() as session:
            lines = (
                session.query(LineListItem)
                .filter(LineListItem.project_id == project_id)
                .order_by(LineListItem.line_number.asc())
                .all()
            )

            wb = Workbook()
            ws = wb.active
            ws.title = "Piping Line List"
            ws.views.sheetView[0].showGridLines = True

            headers = [
                "Line Number", "P&ID No", "Iso No", "Fluid Code", "Fluid Name", "ASME B31.3 Service",
                "Design Press (barg)", "Design Temp (°C)", "Operating Press (barg)", "Operating Temp (°C)",
                "Test Press (barg)", "Test Medium", "Pipe Class", "Size (NPS)", "Schedule", "Material",
                "Insulation", "Tracing", "Painting Code", "RT (%)", "UT (%)", "PWHT Required", "From Point", "To Point", "Status", "Remarks"
            ]

            header_font = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
            header_fill = PatternFill("solid", fgColor="123047")
            center_align = Alignment(horizontal="center", vertical="center")

            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(1, col_idx, h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align

            for row_idx, l in enumerate(lines, 2):
                ws.cell(row_idx, 1, l.line_number)
                ws.cell(row_idx, 2, l.pid_number)
                ws.cell(row_idx, 3, l.iso_number)
                ws.cell(row_idx, 4, l.fluid_code)
                ws.cell(row_idx, 5, l.fluid_name)
                ws.cell(row_idx, 6, l.fluid_service)
                ws.cell(row_idx, 7, l.design_pressure_barg)
                ws.cell(row_idx, 8, l.design_temp_c)
                ws.cell(row_idx, 9, l.operating_pressure_barg)
                ws.cell(row_idx, 10, l.operating_temp_c)
                ws.cell(row_idx, 11, l.test_pressure_barg)
                ws.cell(row_idx, 12, l.test_medium)
                ws.cell(row_idx, 13, l.pipe_class)
                ws.cell(row_idx, 14, l.size_nps)
                ws.cell(row_idx, 15, l.schedule)
                ws.cell(row_idx, 16, l.material)
                ws.cell(row_idx, 17, l.insulation)
                ws.cell(row_idx, 18, l.tracing)
                ws.cell(row_idx, 19, l.painting_code)
                ws.cell(row_idx, 20, l.ndt_percent_rt)
                ws.cell(row_idx, 21, l.ndt_percent_ut)
                ws.cell(row_idx, 22, "YES" if l.pwht_required else "NO")
                ws.cell(row_idx, 23, l.from_point)
                ws.cell(row_idx, 24, l.to_point)
                ws.cell(row_idx, 25, l.status)
                ws.cell(row_idx, 26, l.remarks)

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

            wb.save(path_obj)
            wb.close()
            return str(path_obj)